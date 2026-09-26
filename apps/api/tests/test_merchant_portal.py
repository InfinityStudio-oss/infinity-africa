"""Self-service /v1/merchant/* API: own-merchant resolution, role
enforcement, the overview aggregation, withdrawals dispatch, and
transaction-by-reference lookup — end to end against the in-memory
FakeSupabaseClient (see tests/fakes.py), same pattern as test_routers.py.
"""

import io
import uuid
from decimal import Decimal

import pytest
import resend
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.services.ledger import post_collection_entries, post_disbursement_entries
from app.services.selcom.client import get_selcom_client
from tests.factories import (
    TEST_JWT_SECRET,
    auth_headers,
    create_merchant,
    create_pricing_rule,
    make_merchant_member,
    make_super_admin,
)

client = TestClient(app)


@pytest.fixture(autouse=True)
def _configure_settings(monkeypatch):
    monkeypatch.setenv("SUPABASE_JWT_SECRET", TEST_JWT_SECRET)
    monkeypatch.setenv("MOCK_PROVIDER_FAILURE_RATE", "0")
    monkeypatch.setenv("MOCK_PROVIDER_LATENCY_SECONDS", "0")
    monkeypatch.setenv("DISBURSEMENT_APPROVAL_THRESHOLD", "1000000")
    monkeypatch.setenv("RESEND_API_KEY", "test-resend-key-do-not-use-in-production")
    get_settings.cache_clear()
    get_selcom_client.cache_clear()
    yield
    get_settings.cache_clear()
    get_selcom_client.cache_clear()


class _FakeResend:
    """See tests/test_invoices.py's identical fixture for why this patches
    resend.Emails.send directly rather than going through a client
    abstraction — app/services/email.py calls it as a bare classmethod."""

    def __init__(self):
        self.calls: list[dict] = []
        self.should_fail = False

    def send(self, params: dict) -> dict:
        self.calls.append(params)
        if self.should_fail:
            raise Exception("Resend rejected the request")  # noqa: TRY002
        return {"id": "resend-test-message-id"}


@pytest.fixture(autouse=True)
def fake_resend(monkeypatch):
    fake = _FakeResend()
    monkeypatch.setattr(resend.Emails, "send", fake.send)
    return fake


def _request_and_verify_withdrawal(fake_client, user_id, body: dict, *, idem: str | None = None):
    """The withdrawal flow is two steps now: POST /withdrawals returns an OTP
    challenge and creates nothing, then verifying the emailed code creates
    the PENDING_ADMIN_APPROVAL disbursement. Reads the code straight out of
    the challenge row — it is only ever stored hashed, so the test hashes
    candidate codes rather than reading a plaintext that does not exist."""
    import hashlib

    response = client.post(
        "/v1/merchant/withdrawals",
        headers={**auth_headers(user_id), "Idempotency-Key": idem or _idem()},
        json=body,
    )
    if response.status_code != 202:
        return response
    challenge_id = response.json()["data"]["challenge_id"]
    row = next(
        r for r in fake_client.table("merchant_withdrawal_otp_challenges")._table.rows if r["id"] == challenge_id
    )
    code = next(
        f"{n:06d}" for n in range(1_000_000)
        if hashlib.sha256(f"{n:06d}".encode()).hexdigest() == row["otp_hash"]
    )
    return client.post(
        f"/v1/merchant/withdrawals/{challenge_id}/verify",
        headers=auth_headers(user_id),
        json={"code": code},
    )


def _merchant_and_member(fake_client, role: str = "MERCHANT_ADMIN", **merchant_overrides):
    merchant = create_merchant(fake_client, **merchant_overrides)
    merchant_id = uuid.UUID(merchant["id"])
    user_id = uuid.uuid4()
    make_merchant_member(fake_client, merchant_id, user_id, role)
    return merchant_id, user_id


def _fund_wallet(fake_client, merchant_id: uuid.UUID, amount: str, currency: str = "TZS") -> None:
    fake_client.seed(
        "ledger_accounts",
        {
            "merchant_id": str(merchant_id),
            "name": "Merchant Wallet (test)",
            "account_type": "liability",
            "purpose": "merchant_wallet",
            "currency": currency,
            "balance": amount,
        },
    )


def _idem() -> str:
    return uuid.uuid4().hex


# --- Own-merchant resolution -------------------------------------------------


def test_me_404s_with_no_membership(fake_client):
    user_id = uuid.uuid4()
    response = client.get("/v1/merchant/me", headers=auth_headers(user_id))
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "not_found"


def test_me_404s_with_inactive_membership(fake_client):
    merchant = create_merchant(fake_client)
    user_id = uuid.uuid4()
    fake_client.seed(
        "merchant_users",
        {"merchant_id": merchant["id"], "user_id": str(user_id), "role": "MERCHANT_ADMIN", "status": "invited"},
    )
    response = client.get("/v1/merchant/me", headers=auth_headers(user_id))
    assert response.status_code == 404


def test_me_returns_own_merchant_for_valid_membership(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client)
    response = client.get("/v1/merchant/me", headers=auth_headers(user_id))
    assert response.status_code == 200
    assert response.json()["data"]["id"] == str(merchant_id)


def test_super_admin_has_no_own_merchant(fake_client):
    """/v1/merchant/* is a merchant-dashboard-only surface — a super admin
    with no merchant_users row of their own gets the same 404 anyone else
    without a membership would, no bypass."""
    admin_id = uuid.uuid4()
    make_super_admin(fake_client, admin_id)
    response = client.get("/v1/merchant/me", headers=auth_headers(admin_id))
    assert response.status_code == 404


# --- Role enforcement ---------------------------------------------------------


def test_staff_can_create_payment_link(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client, role="MERCHANT_STAFF")
    response = client.post(
        "/v1/merchant/payment-links",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000", "currency": "TZS"},
    )
    assert response.status_code == 201
    assert response.json()["data"]["merchant_id"] == str(merchant_id)


def test_developer_cannot_create_payment_link(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/payment-links",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000", "currency": "TZS"},
    )
    assert response.status_code == 403


def test_create_payment_link_is_rate_limited(fake_client):
    """Wiring test for app/core/rate_limit.py's payment_link_create scope
    (added during the MVP-readiness rate-limiting sweep) — pre-fills the
    shared in-memory bucket directly rather than firing 30 real requests,
    then confirms one more real request is actually rejected."""
    from app.core.rate_limit import _limiter

    _merchant_id, user_id = _merchant_and_member(fake_client)
    for _ in range(30):
        _limiter.check("payment_link_create:testclient", limit=30, window_seconds=60)

    response = client.post(
        "/v1/merchant/payment-links",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000", "currency": "TZS"},
    )
    assert response.status_code == 429


# --- Payment link extended fields / edit route -------------------------------


def test_create_payment_link_accepts_new_fields(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client)
    response = client.post(
        "/v1/merchant/payment-links",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={
            "amount": "1000",
            "currency": "TZS",
            "customer_email": "grace@example.com",
            "merchant_reference": "ORDER-4821",
            "success_redirect_url": "https://merchant.example.com/thank-you",
            "failure_redirect_url": "https://merchant.example.com/try-again",
        },
    )
    assert response.status_code == 201, response.text
    body = response.json()["data"]
    assert body["customer_email"] == "grace@example.com"
    assert body["merchant_reference"] == "ORDER-4821"
    assert body["success_redirect_url"] == "https://merchant.example.com/thank-you"
    assert body["failure_redirect_url"] == "https://merchant.example.com/try-again"
    assert body["paid_at"] is None
    assert body["attempt_count"] == 0


def test_update_payment_link_while_active_succeeds(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client)
    link = client.post(
        "/v1/merchant/payment-links",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000", "currency": "TZS", "description": "Original"},
    ).json()["data"]

    response = client.patch(
        f"/v1/merchant/payment-links/{link['id']}",
        headers=auth_headers(user_id),
        json={"description": "Updated description", "amount": "2000"},
    )

    assert response.status_code == 200, response.text
    body = response.json()["data"]
    assert body["description"] == "Updated description"
    assert Decimal(body["amount"]) == Decimal(2000)


def test_update_payment_link_rejects_non_active(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client)
    link = client.post(
        "/v1/merchant/payment-links",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000", "currency": "TZS"},
    ).json()["data"]
    client.patch(f"/v1/merchant/payment-links/{link['id']}/cancel", headers=auth_headers(user_id))

    response = client.patch(
        f"/v1/merchant/payment-links/{link['id']}",
        headers=auth_headers(user_id),
        json={"description": "Should not apply"},
    )

    assert response.status_code == 409


def test_update_payment_link_from_another_merchant_not_found(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client)
    link = client.post(
        "/v1/merchant/payment-links",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000", "currency": "TZS"},
    ).json()["data"]

    _other_merchant_id, other_user_id = _merchant_and_member(fake_client, contact_email="other@example.com")
    response = client.patch(
        f"/v1/merchant/payment-links/{link['id']}",
        headers=auth_headers(other_user_id),
        json={"description": "Not yours"},
    )

    assert response.status_code == 404


def test_attempt_count_reflects_collections_against_the_link(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client)
    link = client.post(
        "/v1/merchant/payment-links",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000", "currency": "TZS"},
    ).json()["data"]

    fake_client.seed(
        "collections",
        {
            "merchant_id": str(merchant_id),
            "payment_link_id": link["id"],
            "method": "STK_PUSH",
            "amount": "1000.00",
            "currency": "TZS",
            "status": "failed",
        },
    )
    fake_client.seed(
        "collections",
        {
            "merchant_id": str(merchant_id),
            "payment_link_id": link["id"],
            "method": "STK_PUSH",
            "amount": "1000.00",
            "currency": "TZS",
            "status": "successful",
        },
    )

    get_response = client.get(f"/v1/merchant/payment-links/{link['id']}", headers=auth_headers(user_id))
    assert get_response.json()["data"]["attempt_count"] == 2

    list_response = client.get("/v1/merchant/payment-links", headers=auth_headers(user_id))
    listed = next(row for row in list_response.json()["data"] if row["id"] == link["id"])
    assert listed["attempt_count"] == 2


def test_staff_cannot_create_withdrawal(fake_client):
    """The one deliberate tightening vs. the existing /v1/disbursements/*
    routes, which also allow MERCHANT_STAFF."""
    merchant_id, user_id = _merchant_and_member(fake_client, role="MERCHANT_STAFF")
    _fund_wallet(fake_client, merchant_id, "500000")
    response = client.post(
        "/v1/merchant/withdrawals",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={
            "method": "MOBILE_MONEY",
            "amount": "10000",
            "destination_name": "Jane",
            "destination_phone": "+255700000000",
            "destination_code": "MPESA",
        },
    )
    assert response.status_code == 403


def test_staff_can_view_withdrawals(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="MERCHANT_STAFF")
    response = client.get("/v1/merchant/withdrawals", headers=auth_headers(user_id))
    assert response.status_code == 200


# --- Self-service collections (/v1/merchant/collections/*) -----------------


def test_staff_can_initiate_stk_push_collection(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client, role="MERCHANT_STAFF")
    response = client.post(
        "/v1/merchant/collections/stk-push",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000.00", "customer_phone": "+255700000000"},
    )
    assert response.status_code == 202, response.text
    data = response.json()["data"]
    assert data["merchant_id"] == str(merchant_id)
    assert data["status"] == "processing"
    assert data["transaction_reference"]
    assert data["message"]


def test_stk_push_collection_is_rate_limited(fake_client):
    """Wiring test for app/core/rate_limit.py's merchant_collection_create
    scope (MVP security-hardening pass — the six merchant-portal
    Request Collection push/QR/hosted-checkout/wallet-push endpoints, plus
    create-order-minimal, previously had no rate limit at all, unlike their
    /v1/collections/{method} API-key equivalents in collections_api.py).
    Same pre-fill-the-bucket pattern as
    test_create_payment_link_is_rate_limited above — one endpoint stands
    in for all seven, which share the same scope."""
    from app.core.rate_limit import _limiter

    _merchant_id, user_id = _merchant_and_member(fake_client, role="MERCHANT_STAFF")
    for _ in range(20):
        _limiter.check("merchant_collection_create:testclient", limit=20, window_seconds=60)

    response = client.post(
        "/v1/merchant/collections/stk-push",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000.00", "customer_phone": "+255700000000"},
    )
    assert response.status_code == 429


def test_self_service_dynamic_qr_returns_qr_payload(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="MERCHANT_ADMIN")
    response = client.post(
        "/v1/merchant/collections/dynamic-qr",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "2500.00"},
    )
    assert response.status_code == 202, response.text
    data = response.json()["data"]
    assert data["qr_payload"]
    assert data["expires_at"] == data["qr_expires_at"]


def test_self_service_collection_rejects_expired_payment_link(fake_client):
    """Regression test: /v1/merchant/collections/* previously skipped
    payment_link_id validation entirely (unlike the flat /v1/collections/*
    router), so a merchant's own dashboard/API-key flow could silently
    initiate a collection against an already-expired payment link."""
    _merchant_id, user_id = _merchant_and_member(fake_client, role="MERCHANT_ADMIN")
    link = client.post(
        "/v1/merchant/payment-links",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000", "currency": "TZS"},
    ).json()["data"]
    fake_client.table("payment_links")._table.rows[0]["status"] = "EXPIRED"
    assert fake_client.table("payment_links")._table.rows[0]["id"] == link["id"]

    response = client.post(
        "/v1/merchant/collections/stk-push",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"amount": "1000.00", "customer_phone": "+255700000000", "payment_link_id": link["id"]},
    )

    assert response.status_code == 409
    assert len(fake_client.table("collections")._table.rows) == 0


def test_developer_can_create_api_key(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Dev key", "scopes": ["collections:write", "collections:read"]},
    )
    assert response.status_code == 201
    body = response.json()["data"]
    assert "plaintext_key" in body
    # sk_test_ for sandbox, sk_live_ for live — see _generate_api_key.
    assert body["plaintext_key"].startswith("sk_test_")
    assert body["public_key"].startswith("pk_test_")
    assert body["scopes"] == ["collections:write", "collections:read"]


def test_create_api_key_blocked_when_merchant_api_keys_disabled(fake_client, monkeypatch):
    """ENABLE_MERCHANT_API_KEYS=false (app/core/feature_flags.py) must
    block new key creation/rotation without touching any existing issued
    key — see require_merchant_api_keys_enabled's own docstring."""
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    monkeypatch.setenv("ENABLE_MERCHANT_API_KEYS", "false")
    get_settings.cache_clear()

    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Dev key", "scopes": ["collections:write"]},
    )

    assert response.status_code == 503, response.text
    assert response.json()["error"]["code"] == "feature_disabled"
    assert fake_client.table("api_keys")._table.rows == []


def test_create_api_key_rejects_unknown_scope(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Bad key", "scopes": ["not_a_real_scope"]},
    )
    assert response.status_code == 422


def test_create_api_key_defaults_to_continue_without_ip_whitelist_and_returns_key_last4(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Dev key", "scopes": ["collections:write"]},
    )
    assert response.status_code == 201, response.text
    body = response.json()["data"]
    assert body["ip_whitelist_enabled"] is False
    assert body["continue_without_ip_whitelist"] is True
    assert body["key_last4"] == body["plaintext_key"][-4:]


def test_create_api_key_reconciles_conflicting_ip_whitelist_flags(fake_client):
    """ip_whitelist_enabled wins if a client naively sends both true."""
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={
            "name": "Dev key",
            "scopes": ["collections:write"],
            "ip_whitelist_enabled": True,
            "continue_without_ip_whitelist": True,
            "allowed_ips": [{"ip_address_or_cidr": "41.59.10.20"}],
        },
    )
    assert response.status_code == 201, response.text
    body = response.json()["data"]
    assert body["ip_whitelist_enabled"] is True
    assert body["continue_without_ip_whitelist"] is False


def test_create_api_key_stores_a_label_when_the_merchant_leaves_it_blank(fake_client):
    """api_ip_allowlist.label is NOT NULL, but the request schema leaves the
    label optional -- the portal's inline "Allowed server IPs" list lets you
    add an IP without one. Passing that through as null reached PostgREST as
    a null and 500d, which the portal showed as "Couldn't reach InfinityPay".
    The fake client does not enforce NOT NULL, so this asserts the stored
    value directly rather than relying on an insert failing."""
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")

    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={
            "name": "Website checkout API",
            "scopes": ["collections:write"],
            "ip_whitelist_enabled": True,
            "allowed_ips": [{"ip_address_or_cidr": "41.86.0.0/24"}],
        },
    )

    assert response.status_code == 201, response.text
    rows = fake_client.table("api_ip_allowlist")._table.rows
    assert len(rows) == 1
    # Falls back to the IP itself, as ApiKeyAllowedIp documents.
    assert rows[0]["label"] == "41.86.0.0/24"
    assert rows[0]["status"] == "pending"


def test_create_api_key_keeps_a_label_the_merchant_did_supply(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")

    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={
            "name": "Website checkout API",
            "scopes": ["collections:write"],
            "ip_whitelist_enabled": True,
            "allowed_ips": [{"ip_address_or_cidr": "41.86.0.0/24", "label": "Main ecommerce server"}],
        },
    )

    assert response.status_code == 201, response.text
    assert fake_client.table("api_ip_allowlist")._table.rows[0]["label"] == "Main ecommerce server"


def test_create_api_key_with_ip_whitelisting_enabled_requires_at_least_one_ip(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Dev key", "scopes": ["collections:write"], "ip_whitelist_enabled": True},
    )
    assert response.status_code == 422, response.text
    assert "Add at least one allowed server IP" in response.text


def test_create_api_key_with_ip_whitelisting_enabled_rejects_duplicate_ips(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={
            "name": "Dev key",
            "scopes": ["collections:write"],
            "ip_whitelist_enabled": True,
            "allowed_ips": [
                {"ip_address_or_cidr": "41.59.10.20"},
                {"ip_address_or_cidr": "41.59.10.20", "label": "Duplicate"},
            ],
        },
    )
    assert response.status_code == 422, response.text
    assert "Duplicate" in response.text


def test_create_api_key_with_ip_whitelisting_enabled_rejects_invalid_ip(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={
            "name": "Dev key",
            "scopes": ["collections:write"],
            "ip_whitelist_enabled": True,
            "allowed_ips": [{"ip_address_or_cidr": "not-an-ip"}],
        },
    )
    assert response.status_code == 422, response.text


def test_create_api_key_accepts_a_cidr_and_creates_linked_pending_allowlist_entries(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={
            "name": "Dev key",
            "scopes": ["collections:write"],
            "ip_whitelist_enabled": True,
            "allowed_ips": [
                {"ip_address_or_cidr": "41.59.10.20/32", "label": "Main ecommerce server"},
                {"ip_address_or_cidr": "41.59.11.0/24"},
            ],
        },
    )
    assert response.status_code == 201, response.text
    key_id = response.json()["data"]["id"]

    entries = [
        row
        for row in fake_client.table("api_ip_allowlist")._table.rows
        if row["api_key_id"] == key_id
    ]
    assert len(entries) == 2
    assert {e["ip_address_or_cidr"] for e in entries} == {"41.59.10.20/32", "41.59.11.0/24"}
    assert all(e["status"] == "pending" for e in entries)
    assert all(e["merchant_id"] == str(merchant_id) for e in entries)
    labeled = next(e for e in entries if e["ip_address_or_cidr"] == "41.59.10.20/32")
    assert labeled["label"] == "Main ecommerce server"


def test_create_api_key_without_ip_whitelisting_ignores_any_submitted_ips(fake_client):
    """continue_without_ip_whitelist wipes allowed_ips server-side even if
    the client sent some — the choice is binary, not additive."""
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={
            "name": "Dev key",
            "scopes": ["collections:write"],
            "allowed_ips": [{"ip_address_or_cidr": "41.59.10.20"}],
        },
    )
    assert response.status_code == 201, response.text
    key_id = response.json()["data"]["id"]
    entries = [row for row in fake_client.table("api_ip_allowlist")._table.rows if row["api_key_id"] == key_id]
    assert entries == []


def test_create_api_key_writes_an_audit_log_with_ip_and_user_agent(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    response = client.post(
        "/v1/merchant/api-keys",
        headers={**auth_headers(user_id), "User-Agent": "test-integration/1.0"},
        json={"name": "Dev key", "scopes": ["collections:write"]},
    )
    assert response.status_code == 201, response.text

    logs = [row for row in fake_client.table("audit_logs")._table.rows if row["action"] == "api_key.created"]
    assert len(logs) == 1
    assert logs[0]["user_agent"] == "test-integration/1.0"
    assert logs[0]["ip_address"] is not None


def test_rename_api_key(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    created = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Old name", "scopes": ["collections:write"]},
    ).json()["data"]

    response = client.patch(
        f"/v1/merchant/api-keys/{created['id']}",
        headers=auth_headers(user_id),
        json={"name": "New name"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["data"]["name"] == "New name"


def test_list_api_keys_never_includes_plaintext_or_hash(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Dev key", "scopes": ["collections:write"]},
    )

    response = client.get("/v1/merchant/api-keys", headers=auth_headers(user_id))

    assert response.status_code == 200
    key = response.json()["data"][0]
    assert "plaintext_key" not in key
    assert "hashed_key" not in key
    assert key["scopes"] == ["collections:write"]


def test_revoke_api_key_keeps_scopes_and_marks_revoked(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    created = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Dev key", "scopes": ["invoices:read"]},
    ).json()["data"]

    response = client.patch(
        f"/v1/merchant/api-keys/{created['id']}/revoke", headers=auth_headers(user_id)
    )

    assert response.status_code == 200, response.text
    body = response.json()["data"]
    assert body["status"] == "revoked"
    assert body["scopes"] == ["invoices:read"]


def test_rotate_api_key_revokes_old_and_creates_new_with_same_settings(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    create_pricing_rule(fake_client, merchant_id=merchant_id)
    created = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Prod key", "environment": "live", "scopes": ["collections:write", "collections:read"]},
    ).json()["data"]

    response = client.post(f"/v1/merchant/api-keys/{created['id']}/rotate", headers=auth_headers(user_id))

    assert response.status_code == 201, response.text
    new_key = response.json()["data"]
    assert new_key["id"] != created["id"]
    assert new_key["name"] == "Prod key"
    assert new_key["environment"] == "live"
    assert new_key["scopes"] == ["collections:write", "collections:read"]
    assert "plaintext_key" in new_key
    assert new_key["plaintext_key"].startswith("sk_live_")
    assert new_key["public_key"].startswith("pk_live_")

    rows = {row["id"]: row for row in fake_client.table("api_keys")._table.rows}
    assert rows[created["id"]]["status"] == "revoked"
    assert rows[new_key["id"]]["status"] == "active"


def test_rotate_carries_forward_linked_ip_allowlist_entries_to_the_new_key(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    create_pricing_rule(fake_client, merchant_id=merchant_id)
    created = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={
            "name": "Prod key",
            "environment": "live",
            "scopes": ["collections:write"],
            "ip_whitelist_enabled": True,
            "allowed_ips": [{"ip_address_or_cidr": "41.59.10.20"}],
        },
    ).json()["data"]

    response = client.post(f"/v1/merchant/api-keys/{created['id']}/rotate", headers=auth_headers(user_id))
    assert response.status_code == 201, response.text
    new_key_id = response.json()["data"]["id"]

    new_entries = [
        row for row in fake_client.table("api_ip_allowlist")._table.rows if row["api_key_id"] == new_key_id
    ]
    assert len(new_entries) == 1
    assert new_entries[0]["ip_address_or_cidr"] == "41.59.10.20"
    assert new_entries[0]["status"] == "pending"

    old_entries = [
        row for row in fake_client.table("api_ip_allowlist")._table.rows if row["api_key_id"] == created["id"]
    ]
    assert len(old_entries) == 1  # left alone, historical record tied to the now-revoked key


def test_toggle_ip_whitelist_requires_a_linked_ip_before_enabling(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    created = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Key", "scopes": ["collections:write"]},
    ).json()["data"]

    response = client.patch(
        f"/v1/merchant/api-keys/{created['id']}/ip-whitelist",
        headers=auth_headers(user_id),
        json={"ip_whitelist_enabled": True},
    )
    assert response.status_code == 422, response.text


def test_toggle_ip_whitelist_succeeds_once_a_linked_ip_exists(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    created = client.post(
        "/v1/merchant/api-keys",
        headers=auth_headers(user_id),
        json={"name": "Key", "scopes": ["collections:write"]},
    ).json()["data"]
    client.post(
        "/v1/merchant/ip-allowlist",
        headers=auth_headers(user_id),
        json={
            "environment": "sandbox",
            "label": "Office",
            "ip_address_or_cidr": "41.59.10.20",
            "api_key_id": created["id"],
        },
    )

    response = client.patch(
        f"/v1/merchant/api-keys/{created['id']}/ip-whitelist",
        headers=auth_headers(user_id),
        json={"ip_whitelist_enabled": True},
    )
    assert response.status_code == 200, response.text
    body = response.json()["data"]
    assert body["ip_whitelist_enabled"] is True
    assert body["continue_without_ip_whitelist"] is False

    back_to_open = client.patch(
        f"/v1/merchant/api-keys/{created['id']}/ip-whitelist",
        headers=auth_headers(user_id),
        json={"ip_whitelist_enabled": False},
    )
    assert back_to_open.json()["data"]["continue_without_ip_whitelist"] is True


def test_rotate_already_revoked_key_is_rejected(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="DEVELOPER")
    created = client.post(
        "/v1/merchant/api-keys", headers=auth_headers(user_id), json={"name": "Key", "scopes": ["invoices:read"]}
    ).json()["data"]
    client.patch(f"/v1/merchant/api-keys/{created['id']}/revoke", headers=auth_headers(user_id))

    response = client.post(f"/v1/merchant/api-keys/{created['id']}/rotate", headers=auth_headers(user_id))
    assert response.status_code == 409, response.text


def test_staff_cannot_manage_api_keys(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client, role="MERCHANT_STAFF")
    response = client.post(
        "/v1/merchant/api-keys", headers=auth_headers(user_id), json={"name": "Staff key"}
    )
    assert response.status_code == 403


# --- Invoices: create -> send ------------------------------------------------


def test_create_then_send_invoice(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client)
    create_response = client.post(
        "/v1/merchant/invoices",
        headers=auth_headers(user_id),
        json={
            "customer_email": "customer@example.com",
            "due_date": "2026-12-01",
            "items": [{"description": "Consulting", "quantity": "1", "unit_price": "5000"}],
        },
    )
    assert create_response.status_code == 201
    invoice = create_response.json()["data"]
    assert invoice["status"] == "DRAFT"

    send_response = client.post(
        f"/v1/merchant/invoices/{invoice['id']}/send", headers=auth_headers(user_id)
    )
    assert send_response.status_code == 200
    assert send_response.json()["data"]["status"] == "SENT"


# --- Overview aggregation ------------------------------------------------------


def test_overview_aggregates_all_fields(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client)

    # Two successful collections: 10,000 gross/300 fee, and 5,000 gross/150 fee.
    fake_client.seed(
        "transactions",
        {
            "merchant_id": str(merchant_id),
            "reference": "TXN-1",
            "type": "collection",
            "method": "USSD_PUSH",
            "gross_amount": "10000",
            "fee_amount": "300",
            "net_amount": "9700",
            "currency": "TZS",
            "status": "successful",
        },
    )
    fake_client.seed(
        "transactions",
        {
            "merchant_id": str(merchant_id),
            "reference": "TXN-2",
            "type": "collection",
            "method": "STK_PUSH",
            "gross_amount": "5000",
            "fee_amount": "150",
            "net_amount": "4850",
            "currency": "TZS",
            "status": "successful",
        },
    )
    # One pending, one processing transaction (of any type).
    fake_client.seed(
        "transactions",
        {
            "merchant_id": str(merchant_id),
            "reference": "TXN-3",
            "type": "collection",
            "method": "STK_PUSH",
            "gross_amount": "2000",
            "fee_amount": "0",
            "net_amount": "2000",
            "currency": "TZS",
            "status": "pending",
        },
    )
    fake_client.seed(
        "transactions",
        {
            "merchant_id": str(merchant_id),
            "reference": "TXN-4",
            "type": "disbursement",
            "method": "MOBILE_MONEY",
            "gross_amount": "1000",
            "fee_amount": "0",
            "net_amount": "1000",
            "currency": "TZS",
            "status": "processing",
        },
    )
    # One successful disbursement.
    fake_client.seed(
        "transactions",
        {
            "merchant_id": str(merchant_id),
            "reference": "TXN-5",
            "type": "disbursement",
            "method": "SELCOM_PESA",
            "gross_amount": "3000",
            "fee_amount": "0",
            "net_amount": "3000",
            "currency": "TZS",
            "status": "successful",
        },
    )

    # Payment links: one active-not-expired, one active-but-expired, one cancelled.
    fake_client.seed(
        "payment_links",
        {
            "merchant_id": str(merchant_id),
            "amount": "1000",
            "currency": "TZS",
            "allowed_payment_methods": ["USSD_PUSH"],
            "public_slug": "slug-active",
            "status": "ACTIVE",
            "expires_at": None,
        },
    )
    fake_client.seed(
        "payment_links",
        {
            "merchant_id": str(merchant_id),
            "amount": "1000",
            "currency": "TZS",
            "allowed_payment_methods": ["USSD_PUSH"],
            "public_slug": "slug-expired",
            "status": "ACTIVE",
            "expires_at": "2020-01-01T00:00:00+00:00",
        },
    )
    fake_client.seed(
        "payment_links",
        {
            "merchant_id": str(merchant_id),
            "amount": "1000",
            "currency": "TZS",
            "allowed_payment_methods": ["USSD_PUSH"],
            "public_slug": "slug-cancelled",
            "status": "CANCELLED",
            "expires_at": None,
        },
    )

    # Invoices: one SENT (unpaid), one PAID, one DRAFT.
    fake_client.seed(
        "invoices",
        {
            "merchant_id": str(merchant_id),
            "invoice_number": "INV-1",
            "due_date": "2026-12-01",
            "currency": "TZS",
            "subtotal": "1000",
            "tax_amount": "0",
            "discount_amount": "0",
            "total_amount": "1000",
            "amount_paid": "0",
            "status": "SENT",
        },
    )
    fake_client.seed(
        "invoices",
        {
            "merchant_id": str(merchant_id),
            "invoice_number": "INV-2",
            "due_date": "2026-12-01",
            "currency": "TZS",
            "subtotal": "1000",
            "tax_amount": "0",
            "discount_amount": "0",
            "total_amount": "1000",
            "amount_paid": "1000",
            "status": "PAID",
        },
    )
    fake_client.seed(
        "invoices",
        {
            "merchant_id": str(merchant_id),
            "invoice_number": "INV-3",
            "due_date": "2026-12-01",
            "currency": "TZS",
            "subtotal": "1000",
            "tax_amount": "0",
            "discount_amount": "0",
            "total_amount": "1000",
            "amount_paid": "0",
            "status": "DRAFT",
        },
    )

    _fund_wallet(fake_client, merchant_id, "42000")

    response = client.get("/v1/merchant/overview", headers=auth_headers(user_id))
    assert response.status_code == 200
    data = response.json()["data"]

    assert data["merchant"]["id"] == str(merchant_id)
    assert Decimal(data["total_collections"]) == Decimal(15000)
    assert Decimal(data["available_balance"]) == Decimal(42000)
    assert data["pending_transactions"] == 2  # TXN-3 (pending) + TXN-4 (processing)
    assert data["successful_withdrawals"] == 1  # TXN-5
    assert data["active_payment_links"] == 1  # only slug-active
    assert data["unpaid_invoices"] == 1  # only INV-1
    assert Decimal(data["total_fees_charged"]) == Decimal(450)


def test_overview_404s_with_no_membership(fake_client):
    user_id = uuid.uuid4()
    response = client.get("/v1/merchant/overview", headers=auth_headers(user_id))
    assert response.status_code == 404


# --- Wallet ---------------------------------------------------------------


def test_wallet_ledger_returns_entries_newest_first_with_running_balance(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client)

    post_collection_entries(
        fake_client,
        transaction_id=uuid.uuid4(),
        merchant_id=merchant_id,
        gross_amount=Decimal("1000.00"),
        fee_amount=Decimal(0),
        net_amount=Decimal("1000.00"),
        currency="TZS",
    )
    post_disbursement_entries(
        fake_client,
        transaction_id=uuid.uuid4(),
        merchant_id=merchant_id,
        amount=Decimal("300.00"),
        currency="TZS",
    )

    response = client.get("/v1/merchant/wallet/ledger", headers=auth_headers(user_id))
    assert response.status_code == 200, response.text
    body = response.json()
    rows = body["data"]

    assert len(rows) == 2
    assert rows[0]["direction"] == "debit"
    assert Decimal(rows[0]["balance_after"]) == Decimal("700.00")
    assert rows[1]["direction"] == "credit"
    assert Decimal(rows[1]["balance_after"]) == Decimal("1000.00")
    assert body["meta"]["total"] == 2


def test_wallet_ledger_paginates(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client)
    for _ in range(3):
        post_collection_entries(
            fake_client,
            transaction_id=uuid.uuid4(),
            merchant_id=merchant_id,
            gross_amount=Decimal("100.00"),
            fee_amount=Decimal(0),
            net_amount=Decimal("100.00"),
            currency="TZS",
        )

    response = client.get("/v1/merchant/wallet/ledger?page=1&page_size=2", headers=auth_headers(user_id))
    assert response.status_code == 200, response.text
    body = response.json()
    assert len(body["data"]) == 2
    assert body["meta"]["total"] == 3
    assert body["meta"]["total_pages"] == 2


def test_wallet_ledger_404s_with_no_membership(fake_client):
    user_id = uuid.uuid4()
    response = client.get("/v1/merchant/wallet/ledger", headers=auth_headers(user_id))
    assert response.status_code == 404


def _seed_wallet_entry(fake_client, account_id: str, *, direction: str, amount: str, created_at: str, **overrides):
    row = {
        "ledger_account_id": account_id,
        "direction": direction,
        "amount": amount,
        "currency": "TZS",
        "description": None,
        "created_at": created_at,
        **overrides,
    }
    return fake_client.seed("ledger_entries", row)


def test_wallet_ledger_filters_by_date_range_without_resetting_the_running_balance(fake_client):
    """The date filter must apply AFTER the full-history balance replay —
    an entry inside the window should still show its true opening balance
    (accumulated from everything before the window), not 0."""
    merchant_id, user_id = _merchant_and_member(fake_client)
    _fund_wallet(fake_client, merchant_id, "0")
    account = next(
        a
        for a in fake_client.table("ledger_accounts")._table.rows
        if a["merchant_id"] == str(merchant_id) and a["purpose"] == "merchant_wallet"
    )

    _seed_wallet_entry(fake_client, account["id"], direction="credit", amount="1000", created_at="2026-08-10T09:00:00+00:00")
    _seed_wallet_entry(fake_client, account["id"], direction="credit", amount="500", created_at="2026-08-15T09:00:00+00:00")
    _seed_wallet_entry(fake_client, account["id"], direction="debit", amount="200", created_at="2026-08-20T09:00:00+00:00")

    response = client.get(
        "/v1/merchant/wallet/ledger?start_date=2026-08-15&end_date=2026-08-15",
        headers=auth_headers(user_id),
    )
    assert response.status_code == 200, response.text
    rows = response.json()["data"]
    assert len(rows) == 1
    assert Decimal(rows[0]["balance_before"]) == Decimal("1000.00")
    assert Decimal(rows[0]["balance_after"]) == Decimal("1500.00")


def test_wallet_ledger_cannot_see_another_merchants_entries(fake_client):
    merchant_a_id, user_a = _merchant_and_member(fake_client, business_name="Merchant A")
    merchant_b_id, _user_b = _merchant_and_member(fake_client, business_name="Merchant B")
    _fund_wallet(fake_client, merchant_a_id, "0")
    _fund_wallet(fake_client, merchant_b_id, "0")

    account_a = next(
        a for a in fake_client.table("ledger_accounts")._table.rows if a["merchant_id"] == str(merchant_a_id)
    )
    account_b = next(
        a for a in fake_client.table("ledger_accounts")._table.rows if a["merchant_id"] == str(merchant_b_id)
    )
    _seed_wallet_entry(fake_client, account_a["id"], direction="credit", amount="1000", created_at="2026-08-15T09:00:00+00:00")
    _seed_wallet_entry(fake_client, account_b["id"], direction="credit", amount="9999999", created_at="2026-08-15T09:00:00+00:00")

    # merchant_id is never accepted from the client — there is no query
    # param for it at all — so merchant A's own JWT is the only thing
    # that can ever scope this request.
    response = client.get("/v1/merchant/wallet/ledger", headers=auth_headers(user_a))
    assert response.status_code == 200
    rows = response.json()["data"]
    assert len(rows) == 1
    assert Decimal(rows[0]["amount"]) == Decimal("1000.00")


# --- Wallet ledger Excel export ---------------------------------------------


def test_wallet_ledger_export_requires_start_and_end_date(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client)
    response = client.get("/v1/merchant/wallet/ledger/export", headers=auth_headers(user_id))
    assert response.status_code == 422


def test_wallet_ledger_export_rejects_end_date_before_start_date(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client)
    response = client.get(
        "/v1/merchant/wallet/ledger/export?start_date=2026-08-20&end_date=2026-08-01",
        headers=auth_headers(user_id),
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "validation_error"


def test_wallet_ledger_export_404s_with_no_membership(fake_client):
    user_id = uuid.uuid4()
    response = client.get(
        "/v1/merchant/wallet/ledger/export?start_date=2026-08-01&end_date=2026-08-31",
        headers=auth_headers(user_id),
    )
    assert response.status_code == 404


def test_wallet_ledger_export_returns_xlsx_with_correct_content_type_and_filename(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client, business_name="Masanja Traders", merchant_code="27048391")
    _fund_wallet(fake_client, merchant_id, "0")
    account = next(a for a in fake_client.table("ledger_accounts")._table.rows if a["merchant_id"] == str(merchant_id))
    txn = fake_client.seed(
        "transactions",
        {
            "merchant_id": str(merchant_id),
            "reference": "TXN-EXPORT-1",
            "provider_reference": "SELCOM-9",
            "type": "collection",
            "method": "USSD_PUSH",
            "gross_amount": "1000",
            "fee_amount": "15",
            "net_amount": "985",
            "currency": "TZS",
            "status": "successful",
        },
    )
    _seed_wallet_entry(
        fake_client,
        account["id"],
        direction="credit",
        amount="985",
        created_at="2026-08-15T09:00:00+00:00",
        transaction_id=txn["id"],
    )

    response = client.get(
        "/v1/merchant/wallet/ledger/export?start_date=2026-08-01&end_date=2026-08-31",
        headers=auth_headers(user_id),
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert (
        response.headers["content-disposition"]
        == 'attachment; filename="infinitypay-wallet-ledger-2026-08-01-to-2026-08-31.xlsx"'
    )

    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == [
        "Merchant ID",
        "Business Name",
        "Date",
        "Transaction ID",
        "Type",
        "Direction",
        "Reference",
        "Provider Reference",
        "Payment Method",
        "Opening Balance",
        "Amount",
        "Charge / Fee",
        "Net Amount",
        "Closing Balance",
        "Status",
    ]
    assert ws.freeze_panes == "A2"

    data_row = [cell.value for cell in ws[2]]
    assert data_row[0] == "27048391"
    assert data_row[1] == "Masanja Traders"
    assert data_row[3] == txn["id"]
    assert data_row[4] == "collection"
    assert data_row[5] == "credit"
    assert data_row[6] == "TXN-EXPORT-1"
    assert data_row[7] == "SELCOM-9"
    assert data_row[8] == "USSD_PUSH"
    assert data_row[9] == 0.0  # opening balance
    assert data_row[10] == 985.0  # amount
    assert data_row[11] == 15.0  # charge/fee
    assert data_row[12] == 985.0  # net amount
    assert data_row[13] == 985.0  # closing balance
    assert data_row[14] == "successful"


def test_wallet_ledger_export_respects_the_date_range(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client)
    _fund_wallet(fake_client, merchant_id, "0")
    account = next(a for a in fake_client.table("ledger_accounts")._table.rows if a["merchant_id"] == str(merchant_id))

    _seed_wallet_entry(fake_client, account["id"], direction="credit", amount="1000", created_at="2026-08-05T09:00:00+00:00")
    _seed_wallet_entry(fake_client, account["id"], direction="credit", amount="500", created_at="2026-08-15T09:00:00+00:00")

    response = client.get(
        "/v1/merchant/wallet/ledger/export?start_date=2026-08-10&end_date=2026-08-31",
        headers=auth_headers(user_id),
    )
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.max_row == 2  # header + the one entry inside the range
    assert ws[2][10].value == 500.0  # amount column


def test_wallet_ledger_export_with_no_entries_in_range_still_returns_a_valid_workbook(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client)
    _fund_wallet(fake_client, merchant_id, "0")

    response = client.get(
        "/v1/merchant/wallet/ledger/export?start_date=2026-08-01&end_date=2026-08-31",
        headers=auth_headers(user_id),
    )
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.max_row == 1  # header row only, no data rows
    assert next(cell.value for cell in ws[1]) == "Merchant ID"


# --- Invoice payment links --------------------------------------------------


def test_generate_my_invoice_payment_link_does_not_include_hosted_checkout(fake_client):
    """Regression test: app/routers/merchant_portal.py used to build
    allowed_payment_methods from the full CollectionMethod enum (which now
    includes HOSTED_CHECKOUT). payment_links.allowed_payment_methods has
    its own DB CHECK constraint that was deliberately never widened to
    allow HOSTED_CHECKOUT (see app/schemas/enums.py's
    LEGACY_ALLOWED_PAYMENT_METHODS_DEFAULT) — that mismatch 500ed against
    a real Postgres database on every single "generate Pay Now link"
    click, invisible to this fake client since it doesn't enforce CHECK
    constraints. Assert the exact value so a regression is caught by
    shape, not just by a live-database crash."""
    _merchant_id, admin_id = _merchant_and_member(fake_client)

    create_response = client.post(
        "/v1/merchant/invoices",
        headers=auth_headers(admin_id),
        json={
            "customer_name": "Amina Hassan",
            "customer_phone": "+255700000000",
            "customer_email": "amina@example.com",
            "due_date": "2026-09-01",
            "items": [{"description": "Consulting services", "quantity": "2", "unit_price": "500.00"}],
        },
    )
    assert create_response.status_code == 201, create_response.text
    invoice_id = create_response.json()["data"]["id"]

    send_response = client.post(f"/v1/merchant/invoices/{invoice_id}/send", headers=auth_headers(admin_id))
    assert send_response.status_code == 200, send_response.text

    link_response = client.post(f"/v1/merchant/invoices/{invoice_id}/payment-link", headers=auth_headers(admin_id))
    assert link_response.status_code == 200, link_response.text
    link = link_response.json()["data"]

    link_row = next(r for r in fake_client.table("payment_links")._table.rows if r["id"] == link["id"])
    assert link_row["allowed_payment_methods"] == ["USSD_PUSH", "STK_PUSH", "SELCOM_PESA_PUSH", "DYNAMIC_QR"]
    assert "HOSTED_CHECKOUT" not in link_row["allowed_payment_methods"]


def _create_my_invoice(admin_id: uuid.UUID, **overrides) -> dict:
    body = {
        "customer_name": "Amina Hassan",
        "customer_phone": "+255700000000",
        "customer_email": "amina@example.com",
        "due_date": "2026-09-01",
        "items": [{"description": "Consulting services", "quantity": "2", "unit_price": "500.00"}],
        **overrides,
    }
    response = client.post("/v1/merchant/invoices", headers=auth_headers(admin_id), json=body)
    assert response.status_code == 201, response.text
    return response.json()["data"]


def test_my_invoice_send_sends_a_real_email_and_marks_sent(fake_client, fake_resend):
    _merchant_id, admin_id = _merchant_and_member(fake_client)
    invoice = _create_my_invoice(admin_id)

    response = client.post(f"/v1/merchant/invoices/{invoice['id']}/send", headers=auth_headers(admin_id))

    assert response.status_code == 200, response.text
    sent = response.json()["data"]
    assert sent["status"] == "SENT"
    assert sent["sent_at"] is not None
    assert len(fake_resend.calls) == 1
    assert fake_resend.calls[0]["to"] == ["amina@example.com"]


def test_my_invoice_send_requires_customer_email(fake_client, fake_resend):
    _merchant_id, admin_id = _merchant_and_member(fake_client)
    invoice = _create_my_invoice(admin_id, customer_email=None)

    response = client.post(f"/v1/merchant/invoices/{invoice['id']}/send", headers=auth_headers(admin_id))

    assert response.status_code == 422, response.text
    assert "Customer email is required" in response.json()["error"]["message"]
    assert len(fake_resend.calls) == 0


def test_my_invoice_send_stays_draft_when_email_fails(fake_client, fake_resend):
    fake_resend.should_fail = True
    _merchant_id, admin_id = _merchant_and_member(fake_client)
    invoice = _create_my_invoice(admin_id)

    response = client.post(f"/v1/merchant/invoices/{invoice['id']}/send", headers=auth_headers(admin_id))

    assert response.status_code == 502, response.text
    invoice_row = next(r for r in fake_client.table("invoices")._table.rows if r["id"] == invoice["id"])
    assert invoice_row["status"] == "DRAFT"


def test_merchant_cannot_send_another_merchants_invoice(fake_client, fake_resend):
    _merchant_a_id, admin_a = _merchant_and_member(fake_client)
    _merchant_b_id, admin_b = _merchant_and_member(fake_client)
    invoice = _create_my_invoice(admin_a)

    response = client.post(f"/v1/merchant/invoices/{invoice['id']}/send", headers=auth_headers(admin_b))

    assert response.status_code == 404
    assert len(fake_resend.calls) == 0
    invoice_row = next(r for r in fake_client.table("invoices")._table.rows if r["id"] == invoice["id"])
    assert invoice_row["status"] == "DRAFT"


# --- Withdrawals dispatch -------------------------------------------------


@pytest.mark.parametrize(
    ("method", "destination_code", "body_extra"),
    [
        ("SELCOM_PESA", "SELCOM", {"destination_phone": "+255700000001"}),
        ("MOBILE_MONEY", "MPESA", {"destination_phone": "+255700000002"}),
        (
            "BANK_ACCOUNT",
            "CRDB",
            {"bank_name": "CRDB", "bank_account_number": "1234567890", "bank_account_name": "Jane Doe"},
        ),
    ],
)
def test_withdrawal_dispatches_to_correct_method(fake_client, method, destination_code, body_extra):
    merchant_id, user_id = _merchant_and_member(fake_client)
    _fund_wallet(fake_client, merchant_id, "100000")

    response = _request_and_verify_withdrawal(
        fake_client,
        user_id,
        {
            "method": method,
            "destination_code": destination_code,
            "amount": "10000",
            "destination_name": "Jane",
            **body_extra,
        },
    )
    assert response.status_code == 202, response.text
    body = response.json()["data"]
    assert body["method"] == method
    assert body["merchant_id"] == str(merchant_id)
    # Every withdrawal now requires Super Admin approval before Selcom is
    # ever called — never auto-processed, regardless of amount/method.
    assert body["status"] == "PENDING_ADMIN_APPROVAL"


def test_withdrawal_blocked_when_withdrawals_disabled(fake_client, monkeypatch):
    """ENABLE_WITHDRAWALS=false must pause the Merchant Portal withdrawal
    route too — the check runs first, before any withdrawal row is written
    or Selcom is ever considered (Selcom is never called from this path
    regardless)."""
    merchant_id, user_id = _merchant_and_member(fake_client)
    _fund_wallet(fake_client, merchant_id, "100000")
    monkeypatch.setenv("ENABLE_WITHDRAWALS", "false")
    get_settings.cache_clear()

    response = client.post(
        "/v1/merchant/withdrawals",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={
            "method": "SELCOM_PESA",
            "destination_code": "SELCOM",
            "amount": "10000",
            "destination_name": "Masanja",
            "destination_phone": "+255657878545",
        },
    )

    assert response.status_code == 503, response.text
    assert response.json()["error"]["code"] == "feature_disabled"
    assert fake_client.table("disbursements")._table.rows == []


def test_withdrawal_amount_parsed_as_decimal_not_currency_string(fake_client):
    """The API contract is a plain decimal string — a formatted value like
    "TZS 10,000" is rejected as a validation error, never coerced."""
    merchant_id, user_id = _merchant_and_member(fake_client)
    _fund_wallet(fake_client, merchant_id, "100000")

    response = client.post(
        "/v1/merchant/withdrawals",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={
            "method": "SELCOM_PESA",
            "destination_code": "SELCOM",
            "amount": "TZS 10,000",
            "destination_name": "Masanja",
            "destination_phone": "+255657878545",
        },
    )

    assert response.status_code == 422, response.text
    assert fake_client.table("disbursements")._table.rows == []


def test_withdrawal_insufficient_balance_returns_409(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client)
    _fund_wallet(fake_client, merchant_id, "1000")

    response = client.post(
        "/v1/merchant/withdrawals",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={
            "method": "MOBILE_MONEY",
            "amount": "50000",
            "destination_name": "Jane",
            "destination_phone": "+255700000003",
            "destination_code": "MPESA",
        },
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "insufficient_balance"


def test_withdrawal_bank_account_requires_bank_fields(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client)
    response = client.post(
        "/v1/merchant/withdrawals",
        headers={**auth_headers(user_id), "Idempotency-Key": _idem()},
        json={"method": "BANK_ACCOUNT", "amount": "10000", "destination_name": "Jane", "destination_code": "CRDB"},
    )
    assert response.status_code == 422


def test_withdrawal_mobile_money_accepts_optional_network(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client)
    _fund_wallet(fake_client, merchant_id, "100000")

    response = _request_and_verify_withdrawal(
        fake_client,
        user_id,
        {
            "method": "MOBILE_MONEY",
            "amount": "10000",
            "destination_name": "Jane",
            "destination_phone": "+255700000003",
            "destination_code": "MPESA",
            "network": "TIGO_PESA",
        },
    )

    assert response.status_code == 202, response.text
    body = response.json()["data"]
    assert body["status"] == "PENDING_ADMIN_APPROVAL"
    # Not reserved/settled yet — nothing until a Super Admin approves.
    assert body["transaction_reference"] is None
    assert body["total_reserved_amount"] is not None


# --- Transactions by reference -----------------------------------------------


def test_get_transaction_by_reference_found(fake_client):
    merchant_id, user_id = _merchant_and_member(fake_client)
    fake_client.seed(
        "transactions",
        {
            "merchant_id": str(merchant_id),
            "reference": "TXN-FINDME",
            "type": "collection",
            "method": "USSD_PUSH",
            "gross_amount": "1000",
            "fee_amount": "0",
            "net_amount": "1000",
            "currency": "TZS",
            "status": "successful",
            "metadata": {},
        },
    )
    response = client.get("/v1/merchant/transactions/TXN-FINDME", headers=auth_headers(user_id))
    assert response.status_code == 200
    assert response.json()["data"]["reference"] == "TXN-FINDME"


def test_get_transaction_by_reference_not_found(fake_client):
    _merchant_id, user_id = _merchant_and_member(fake_client)
    response = client.get("/v1/merchant/transactions/TXN-NOPE", headers=auth_headers(user_id))
    assert response.status_code == 404


def test_get_transaction_by_reference_cross_tenant_not_leaked(fake_client):
    """Another merchant's transaction, matched by reference alone, must not
    be visible to a caller from a different merchant."""
    other_merchant_id, _other_user = _merchant_and_member(fake_client)
    fake_client.seed(
        "transactions",
        {
            "merchant_id": str(other_merchant_id),
            "reference": "TXN-SHARED-REF",
            "type": "collection",
            "method": "USSD_PUSH",
            "gross_amount": "1000",
            "fee_amount": "0",
            "net_amount": "1000",
            "currency": "TZS",
            "status": "successful",
        },
    )
    _my_merchant_id, my_user_id = _merchant_and_member(fake_client)
    response = client.get("/v1/merchant/transactions/TXN-SHARED-REF", headers=auth_headers(my_user_id))
    assert response.status_code == 404
